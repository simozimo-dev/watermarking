"""
Rigenera i file di esempio in samples/ (non serve per usare i test: i file sono già inclusi).
Richiede: python-docx, openpyxl, python-pptx, reportlab, Pillow, LibreOffice (soffice) e ffmpeg.

    pip install python-docx openpyxl python-pptx reportlab pillow
    python make_samples.py
"""
import email.message
import io
import random
import shutil
import subprocess
import sys
import tempfile
import zipfile
from pathlib import Path

from docx import Document
from docx.shared import Inches
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PInches
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas

OUT = Path(__file__).with_name("samples")
rnd = random.Random(42)
LOREM = ("Lorem ipsum dolor sit amet, consectetur adipiscing elit, sed do eiusmod tempor incididunt ut labore "
         "et dolore magna aliqua. Ut enim ad minim veniam, quis nostrud exercitation ullamco laboris nisi ut "
         "aliquip ex ea commodo consequat. Duis aute irure dolor in reprehenderit in voluptate velit esse.")


def image(path, w, h, fmt):
    img = Image.new("RGB", (w, h), (235, 240, 245))
    d = ImageDraw.Draw(img)
    for _ in range(60):
        x, y = rnd.randrange(w), rnd.randrange(h)
        r = rnd.randrange(20, max(21, w // 6))
        d.ellipse([x - r, y - r, x + r, y + r], fill=tuple(rnd.randrange(60, 220) for _ in range(3)))
    d.text((20, 20), "Sample image for watermark tests", fill=(20, 30, 40))
    img.save(path, fmt, quality=85) if fmt == "JPEG" else img.save(path, fmt)


def memo(path):
    doc = Document()
    doc.add_heading("Internal memo", 1)
    doc.add_paragraph("To: all staff\nFrom: Operations\nSubject: quarterly review")
    for _ in range(4):
        doc.add_paragraph(LOREM)
    doc.save(path)


def report(path, img):
    doc = Document()
    doc.add_heading("Annual report", 0)
    for chapter in range(1, 9):
        doc.add_heading(f"Chapter {chapter}", 1)
        for _ in range(6):
            doc.add_paragraph(LOREM)
        if chapter % 3 == 0:
            doc.add_picture(str(img), width=Inches(5))
        t = doc.add_table(rows=6, cols=4)
        t.style = "Table Grid"
        for r in range(6):
            for c in range(4):
                t.cell(r, c).text = f"{rnd.randint(100, 9999)}"
        doc.add_page_break()
    doc.save(path)


def workbook(path):
    wb = Workbook()
    ws = wb.active
    ws.title = "Budget"
    ws.append(["Department", "Q1", "Q2", "Q3", "Q4"])
    for d in ["Sales", "Marketing", "IT", "HR", "Legal", "Finance", "Operations"]:
        ws.append([d] + [rnd.randint(10_000, 90_000) for _ in range(4)])
    ws2 = wb.create_sheet("Data")
    for i in range(1, 501):
        ws2.append([i, f"Item {i}", rnd.random() * 1000, rnd.choice(["A", "B", "C"])])
    wb.save(path)


def deck(path):
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Quarterly business review"
    title.placeholders[1].text = "Operations team"
    for i in range(1, 6):
        s = prs.slides.add_slide(prs.slide_layouts[1])
        s.shapes.title.text = f"Topic {i}"
        s.placeholders[1].text = "\n".join(LOREM.split(". ")[:4])
    prs.save(path)


def pdf(path, pages):
    c = canvas.Canvas(str(path), pagesize=A4)
    w, h = A4
    for p in range(1, pages + 1):
        c.setFont("Helvetica-Bold", 18)
        c.drawString(60, h - 80, f"Document page {p}")
        c.setFont("Helvetica", 10)
        y = h - 120
        for _ in range(40):
            c.drawString(60, y, LOREM[:95])
            y -= 16
        c.showPage()
    c.save()


def soffice(src, fmt, outdir):
    subprocess.run(["soffice", "--headless", "--convert-to", fmt, "--outdir", str(outdir), str(src)],
                   check=True, capture_output=True, timeout=180)


def main():
    OUT.mkdir(exist_ok=True)
    tmp = Path(tempfile.mkdtemp())
    image(OUT / "scan.png", 1600, 1200, "PNG")
    image(OUT / "photo.jpg", 2400, 1600, "JPEG")
    memo(OUT / "memo.docx")
    report(OUT / "report.docx", OUT / "photo.jpg")
    workbook(OUT / "budget.xlsx")
    deck(OUT / "deck.pptx")
    pdf(OUT / "onepage.pdf", 1)
    pdf(OUT / "manual.pdf", 40)

    soffice(OUT / "memo.docx", "odt", tmp)
    soffice(OUT / "budget.xlsx", "ods", tmp)
    soffice(OUT / "deck.pptx", "odp", tmp)
    soffice(OUT / "memo.docx", "doc", tmp)
    shutil.move(tmp / "memo.odt", OUT / "letter.odt")
    shutil.move(tmp / "budget.ods", OUT / "sheet.ods")
    shutil.move(tmp / "deck.odp", OUT / "slides.odp")
    shutil.move(tmp / "memo.doc", OUT / "legacy.doc")

    subprocess.run(["ffmpeg", "-y", "-loglevel", "error", "-f", "lavfi", "-i", "testsrc=size=640x360:rate=25",
                    "-t", "6", "-pix_fmt", "yuv420p", "-c:v", "libx264", "-preset", "veryfast", "-crf", "32",
                    str(OUT / "clip.mp4")], check=True, timeout=120)

    msg = email.message.EmailMessage()
    msg["From"] = "sender@example.com"
    msg["To"] = "receiver@example.com"
    msg["Subject"] = "Documents for review"
    msg.set_content("Please find the documents attached.\n")
    msg.add_attachment((OUT / "onepage.pdf").read_bytes(), maintype="application", subtype="pdf", filename="onepage.pdf")
    msg.add_attachment((OUT / "memo.docx").read_bytes(), maintype="application",
                       subtype="vnd.openxmlformats-officedocument.wordprocessingml.document", filename="memo.docx")
    (OUT / "message.eml").write_bytes(bytes(msg))

    with zipfile.ZipFile(OUT / "bundle.zip", "w", zipfile.ZIP_DEFLATED) as z:
        for name in ["memo.docx", "onepage.pdf", "scan.png"]:
            z.write(OUT / name, f"docs/{name}")

    shutil.rmtree(tmp, ignore_errors=True)
    for f in sorted(OUT.iterdir()):
        print(f"{f.name:14} {f.stat().st_size:>10,} bytes")


if __name__ == "__main__":
    sys.exit(main())
